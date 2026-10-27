"""The report as `report.docx`, for Google Docs (Drive converts it to a Google Doc on upload) and Word.

Plain document styles, no theme: a document is edited and restyled by whoever receives it. Charts become tables.
The engine never uploads it anywhere (AppScreenMap sends nothing).
"""

from __future__ import annotations

import io
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from .document import MAX_SHOTS
from .evidence import crop
from .markdown import SEV, where_text

THUMB_IN = 0.75  # thumbnail width: a row of seven fits a page
EVIDENCE_IN = 2.0  # evidence crops: the outlined element readable on paper
GALLERY_COLS = 5
# tokens.css --print-accent: the dark accent would be too light on paper.
LINK = RGBColor(0x0B, 0x5C, 0xBA)


def _link(par, url: str, text: str) -> None:
    """python-docx has no hyperlink API: add the relationship and the run by hand."""
    rid = par.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    h = OxmlElement("w:hyperlink")
    h.set(qn("r:id"), rid)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), str(LINK))
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    props.append(color)
    props.append(underline)
    run.append(props)
    t = OxmlElement("w:t")
    t.text = text
    t.set(qn("xml:space"), "preserve")
    run.append(t)
    h.append(run)
    par._p.append(h)


def _labelled(d, label: str, text: str | None) -> None:
    if not text:
        return
    p = d.add_paragraph()
    p.add_run(f"{label}: ").bold = True
    p.add_run(str(text))


def _table(d, head: list[str], rows: list[list], right: set[int] = frozenset()) -> None:
    t = d.add_table(rows=1, cols=len(head))
    t.style = "Light List"
    for i, h in enumerate(head):
        t.rows[0].cells[i].text = h
    for r in rows:
        cells = t.add_row().cells
        for i, v in enumerate(r):
            cells[i].text = "" if v is None else str(v)
            if i in right:
                cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    d.add_paragraph()


def _refs(d, refs: list[dict], label: str = "Sources") -> None:
    if not refs:
        return
    p = d.add_paragraph()
    p.add_run(f"{label}: ").bold = True
    for i, r in enumerate(refs):
        if i:
            p.add_run(" · ")
        _link(p, r["url"], r.get("title") or r["url"])


def _finding(d, f: dict, run: Path, category: str | None) -> None:
    d.add_heading(f"{SEV[f['severity']]} · {f['title']}", level=4)
    meta = [f["id"]] + [x for x in (category, f.get("platform")) if x]
    if f.get("due"):
        meta.append(f"due {f['due']}")
    if f.get("requirement"):
        meta.append(f"requirement: {f['requirement']}")
    if f.get("triage"):
        meta.append(f"triage: {f['triage'].replace('_', ' ')}")
    m = d.add_paragraph()
    r = m.add_run(" · ".join(meta))
    r.font.size, r.font.name = Pt(8), "Courier New"
    if f.get("detail"):
        d.add_paragraph(f["detail"])
    for t in f.get("items") or []:
        d.add_paragraph(t, style="List Bullet")
    _labelled(d, "Where", ", ".join(w["id"] for w in f["where"]) if f.get("items") else where_text(f))
    _labelled(d, "Fix", f.get("fix"))
    _refs(d, f.get("refs") or [])
    where = [w for w in f.get("where") or [] if w.get("image")]
    if where:
        p = d.add_paragraph()
        for w in where[:MAX_SHOTS]:
            if _picture(p, run / w["image"], THUMB_IN):
                p.add_run(" ")
        if len(where) > MAX_SHOTS:
            p.add_run(f"+{len(where) - MAX_SHOTS} more")
    crops = [
        (ev, data) for ev in f.get("evidence") or [] if (data := crop(run, ev.get("shot"), ev.get("bounds")))
    ]
    if crops:
        p = d.add_paragraph()
        p.add_run("Evidence (failing element outlined)").bold = True
        p = d.add_paragraph()
        for ev, data in crops:
            if _picture(p, io.BytesIO(data), EVIDENCE_IN):
                p.add_run("  ")
        d.add_paragraph(" · ".join(f"{ev['id']}: {ev['label']}" for ev, _ in crops)).runs[0].font.size = Pt(8)


def _picture(par, src, width_in: float) -> bool:
    try:
        par.add_run().add_picture(src if isinstance(src, io.BytesIO) else str(src), width=Inches(width_in))
    # python-docx raises a bare Exception on a truncated JPEG: an unreadable image never blocks the report.
    except Exception:  # noqa: BLE001
        return False
    return True


def _gallery(d, gallery: list[dict], run: Path) -> None:
    d.add_heading("Screens", level=1)
    d.add_paragraph(f"Every screen the crawl reached ({len(gallery)}), in map order.")
    rows = [gallery[i : i + GALLERY_COLS] for i in range(0, len(gallery), GALLERY_COLS)]
    t = d.add_table(rows=len(rows), cols=GALLERY_COLS)
    for r, row in enumerate(rows):
        for c, g in enumerate(row):
            cell = t.rows[r].cells[c]
            _picture(cell.paragraphs[0], run / g["image"], 1.1)
            cap = cell.add_paragraph()
            cap.add_run(g["title"]).bold = True
            meta = cap.add_run(f"\n{g['id']}" + (f" · {g['scenario']}" if g.get("scenario") else ""))
            meta.font.size = Pt(8)
    d.add_paragraph()


def render(doc: dict, out_dir: Path) -> Path:
    out_dir = Path(out_dir)
    run = Path(doc["run_dir"])
    c = doc["cover"]
    d = Document()
    d.core_properties.title = doc["title"]
    d.core_properties.author = doc["tool"]

    d.add_heading(c["app"] or "App report", level=0)
    d.add_paragraph(f"AppScreenMap report · profile {doc['profile']}")
    facts = [
        ("Package", c.get("package")),
        ("Version", c.get("version")),
        ("Platform", c.get("platform")),
        ("Canvas", f"{c['canvas']} · {c['screens']} screens"),
        ("Crawled", c.get("crawled")),
        ("Report", f"{doc['generated_at']} · {doc['tool']}"),
    ]
    _table(d, ["", ""], [[k, v] for k, v in facts if v])

    d.add_heading("Summary", level=1)
    _table(
        d,
        ["Audit", "Status", "Score", "Errors", "Warnings", "Info", "Summary"],
        [
            [s["title"], s["status_word"], "—" if s["score"] is None else s["score"]]
            + [s["counts"][k] for k in ("error", "warning", "info")]
            + [s["summary"]]
            for s in doc["summary"]
        ],
        right={2, 3, 4, 5},
    )
    if doc["top"]:
        d.add_heading("Top findings", level=2)
        for f in doc["top"]:
            extra = " · ".join(
                x for x in (f["audit_title"], where_text(f), f"due {f['due']}" if f.get("due") else "") if x
            )
            p = d.add_paragraph(style="List Number")
            p.add_run(f"{SEV[f['severity']]}: ").bold = True
            p.add_run(f"{f['title']} · {extra}")
    if doc["deadlines"]:
        d.add_heading("Deadlines", level=2)
        _table(
            d,
            ["Due", "Finding", "Audit", "State"],
            [[f["due"], f["title"], f["audit_title"], f.get("requirement") or ""] for f in doc["deadlines"]],
        )

    cond = doc["conditions"]
    d.add_heading("Conditions", level=1)
    if cond["device"]:
        d.add_paragraph(f"Captured on {cond['device']}.")
        _table(d, ["", ""], [[k, v] for k, v in cond["rows"]])
    if cond["scenarios"]:
        _table(
            d,
            ["Scenario", "Run", "Screens"],
            [[sc["label"] or sc["id"], sc["run"], sc["screens"]] for sc in cond["scenarios"]],
            right={2},
        )
    if not cond["device"] and not cond["scenarios"]:
        d.add_paragraph("No device facts recorded for this canvas.")

    if doc.get("gallery"):
        _gallery(d, doc["gallery"], run)

    for s in doc["sections"]:
        d.add_heading(s["title"], level=1)
        head = d.add_paragraph()
        head.add_run(s["status_word"]).bold = True
        if s["score"] is not None:
            head.add_run(f" · score {s['score']}")
        if s.get("tool"):
            head.add_run(f" · {s['tool']}")
        d.add_paragraph(s["summary"])
        if s.get("score_note"):
            d.add_paragraph(f"Score: {s['score_note']}").runs[0].italic = True
        if s["metrics"]:
            _table(
                d,
                ["Metric", "Value"],
                [
                    [m["label"], f"{m['value']}{(' ' + m['unit']) if m.get('unit') else ''}"]
                    for m in s["metrics"]
                ],
                right={1},
            )
        for ch in s["charts"]:
            unit = f" ({ch['unit']})" if ch.get("unit") else ""
            d.add_paragraph(ch["title"] + (f" · {ch['note']}" if ch.get("note") else "")).runs[0].bold = True
            rows = [[x["label"], x["value"]] for x in ch["data"]]
            if ch.get("budget") is not None:
                rows.append([ch.get("budgetLabel") or "budget", ch["budget"]])
            _table(d, ["", f"Value{unit}"], rows, right={1})
        if not s["findings"]:
            d.add_paragraph("Not run." if s["status"] == "not_run" else "No findings.")
        for cat, fs in s["groups"].items():
            if len(fs) > 1:
                d.add_heading(cat, level=3)
            for f in fs:
                _finding(d, f, run, None if len(fs) > 1 else cat)
        if s["manual"]:
            d.add_heading("Not tested (manual checks)", level=3)
            for m in s["manual"]:
                p = d.add_paragraph(style="List Bullet")
                p.add_run(f"{m.get('title')}: ").bold = True
                p.add_run(m.get("steps") or m.get("detail") or "")

    ap = doc["appendix"]
    d.add_heading("Appendix", level=1)
    d.add_heading("Method", level=2)
    for m in ap["method"]:
        d.add_paragraph(m, style="List Bullet")
    if ap["sources"]:
        d.add_heading("Sources", level=2)
        for r in ap["sources"]:
            p = d.add_paragraph(style="List Bullet")
            _link(p, r["url"], r.get("title") or r["url"])
            if r.get("publisher"):
                p.add_run(f" ({r['publisher']})")
    if ap["glossary"]:
        d.add_heading("Glossary", level=2)
        _table(d, ["Term", "Meaning"], [[k, v] for k, v in ap["glossary"]])

    out_dir.mkdir(parents=True, exist_ok=True)
    p = out_dir / "report.docx"
    d.save(str(p))
    return p
