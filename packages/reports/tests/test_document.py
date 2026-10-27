from __future__ import annotations

import json
import re

from appscreenmap.reports.document import build_document, write_report
from appscreenmap.reports.fixes import tasks_from


def _screen(sid, title, **kw):
    return {"id": sid, "kind": "app", "title": title, "activity": "dev.alpha.MainActivity", "depth": 0,
            "elements": [], **kw}  # fmt: skip


def _canvas(tmp_path, a11y_issues=None, store=True, extra=None):
    run = tmp_path / "alpha"
    (run / "screens").mkdir(parents=True)
    (run / "screens" / "s001.png").write_bytes(b"\x89PNG fake")
    data = {
        "schema": 1,
        "run": {
            "id": "alpha-1",
            "app": {"package": "dev.alpha", "label": "Alpha", "version": "1.2"},
            "device": {"platform": "android", "name": "Pixel", "android": "17", "sdk": 37},
            "config": {"max_steps": 50},
            "steps": 10,
            "scenarios": [
                {"id": "sc1", "label": "First run", "run_id": "alpha-1", "screens": ["s001", "s002"]}
            ],
        },
        "screens": [_screen("s001", "Home", thumb="screens/s001.png"), _screen("s002", "Profile")],
        "transitions": [],
        "events": [{"step": 4, "type": "crash", "screen": "s002", "detail": "NPE"}],
        "stats": {"screens": 2, "transitions": 1, "max_depth": 1, "elements": {"done": 4, "pending": 1}},
    }
    (run / "map.json").write_text(json.dumps(data))
    (run / "titles.json").write_text(json.dumps({"language": "en", "titles": {"s002": "Your profile"}}))
    rule = {"sc": "1.1.1", "name": "Non-text Content", "level": "A", "severity": "warning", "why": "No label.",
            "fix": "Add a contentDescription."}  # fmt: skip
    issues = a11y_issues if a11y_issues is not None else [
        {"key": f"{s}::image-alt::x", "screen": s, "rule": "image-alt", "sc": "1.1.1", "level": "A",
         "severity": "warning", "rid": "dev.alpha:id/logo", "label": "logo", "triage": "open"}
        for s in ("s001", "s002")
    ]  # fmt: skip
    (run / "a11y.json").write_text(
        json.dumps(
            {
                "rules": {"image-alt": rule},
                "summary": {"by_criterion": {"1.1.1": len(issues)}},
                "issues": issues,
            }
        )
    )
    audits = run / "audits"
    audits.mkdir()
    if store:
        (audits / "store.json").write_text(json.dumps({
            "schema": 1, "audit": "store", "title": "Store readiness", "source": "run", "status": "fail", "score": 40,
            "summary": "Blocked.", "findings": [
                {"id": "store::target", "severity": "error", "title": "Targets API 30", "due": "2026-11-01",
                 "fix": "Raise targetSdk.", "refs": [{"title": "Target API", "url": "https://support.google.com/x"}]},
            ],
        }))  # fmt: skip
    (audits / "platform.json").write_text(json.dumps({
        "schema": 1, "audit": "platform", "title": "Platform changes", "source": "sample", "status": "warn",
        "summary": "Sample.", "findings": [{"id": "p1", "severity": "error", "title": "Sample finding"}],
    }))  # fmt: skip
    (audits / "security-evidence.json").write_text(json.dumps({"schema": 1, "sources": {}}))
    for aid, a in (extra or {}).items():
        (audits / f"{aid}.json").write_text(json.dumps(a))
    return run


def test_report_has_every_audit_in_order_and_skips_sample_data(tmp_path):
    run = _canvas(tmp_path)
    doc = build_document(run)
    assert [s["id"] for s in doc["sections"]] == ["coverage", "stability", "a11y", "perf", "store"]
    assert doc["cover"]["version"] == "1.2"
    # Built-in audits are written as files, like any tool's.
    assert json.loads((run / "audits" / "a11y.json").read_text())["audit"] == "a11y"
    # titles.json (English display titles) wins over map.json's own.
    a11y = next(s for s in doc["sections"] if s["id"] == "a11y")
    assert [w["title"] for w in a11y["findings"][0]["where"]] == ["Home", "Your profile"]
    assert doc["deadlines"][0]["due"] == "2026-11-01"
    assert all(f["audit"] != "coverage" for f in doc["top"])


def test_a_new_audit_file_gets_a_section_without_renderer_changes(tmp_path):
    madeup = {"schema": 1, "audit": "privacy", "title": "Privacy labels", "status": "warn", "score": 70,
              "summary": "One label missing.", "findings": [
                  {"id": "privacy::label", "severity": "warning", "title": "Data safety label missing", "screen": "s001"}]}  # fmt: skip
    run = _canvas(tmp_path, extra={"privacy": madeup})
    out = write_report(run, ["md", "html", "fixes"])
    md, page = out["md"].read_text(), out["html"].read_text()
    assert "## Privacy labels" in md and "Data safety label missing" in md
    assert 'id="privacy"' in page and "Data safety label missing" in page
    assert "privacy::label" in out["fixes"].read_text()


def test_markdown_and_html_keep_the_five_sections_in_order(tmp_path):
    out = write_report(_canvas(tmp_path), ["md", "html"])
    md = out["md"].read_text()
    heads = [h for h in re.findall(r"^## (.+)$", md, re.MULTILINE)]
    assert heads[0] == "Summary" and heads[1] == "Conditions" and heads[-1] == "Appendix"
    assert (out["md"].parent / "report-assets" / "s001.png").is_file()
    page = out["html"].read_text()
    assert (
        page.index('<h2 id="summary"')
        < page.index('<h2 id="conditions"')
        < page.index('id="coverage"')
        < page.index('<h2 id="appendix"')
    )
    assert "data:image/png;base64," in page  # thumbnails inlined: the file opens offline
    assert "--accent" in page  # design tokens inlined
    assert 'font-family:"Geist Variable"' in page  # and the fonts, so it looks the same offline


def test_profiles_and_audit_lists_pick_sections(tmp_path):
    run = _canvas(tmp_path)
    assert [s["id"] for s in build_document(run, profile="release")["sections"]] == ["store"]
    assert [s["id"] for s in build_document(run, only=["a11y", "perf"])["sections"]] == ["a11y", "perf"]


def test_fix_list_groups_orders_and_leaves_out_crawl_and_dismissed(tmp_path):
    doc = build_document(_canvas(tmp_path))
    tasks = tasks_from(doc)
    assert [t["task"] for t in tasks] == [
        "store::target",
        "stability::crash::s002::0",
        "a11y::image-alt",
    ]  # due date first
    a11y = tasks[-1]
    assert [w["id"] for w in a11y["where"]] == ["s001", "s002"]  # one task, every screen
    assert a11y["elements"][0]["rid"] == "dev.alpha:id/logo"

    dismissed = [
        {"key": "k", "screen": "s001", "rule": "image-alt", "sc": "1.1.1", "level": "A", "severity": "warning",
         "triage": "wontfix"}
    ]  # fmt: skip
    doc = build_document(_canvas(tmp_path / "b", a11y_issues=dismissed))
    assert "a11y::image-alt" not in [t["task"] for t in tasks_from(doc)]


def test_fix_list_ticks_a_task_once_its_finding_is_gone(tmp_path):
    run = _canvas(tmp_path)
    first = write_report(run, ["fixes"])["fixes"].read_text()
    assert "- [ ] **1. Targets API 30**" in first
    assert "**Check it's fixed:** Rebuild, then `appscreenmap store`" in first

    store = json.loads((run / "audits" / "store.json").read_text())
    store["findings"], store["status"] = [], "pass"
    (run / "audits" / "store.json").write_text(json.dumps(store))
    second = write_report(run, ["fixes"])["fixes"].read_text()
    assert "- [x] Targets API 30 · `store::target` · gone in `alpha-1`" in second
    assert "- [ ] **1. Targets API 30**" not in second


def test_a_task_whose_audit_did_not_run_again_is_not_ticked(tmp_path):
    run = _canvas(tmp_path)
    write_report(run, ["fixes"])
    (run / "audits" / "store.json").unlink()
    text = write_report(run, ["fixes"])["fixes"].read_text()
    assert "[x] Targets API 30" not in text


def _triage(run, statuses):
    """notes.json triage, as the viewer writes it, keyed like a11y.json issues."""
    (run / "notes.json").write_text(json.dumps({"a11y": {k: {"status": v} for k, v in statuses.items()}}))


def test_triage_in_notes_counts_without_rerunning_a11y_and_is_never_a_fix(tmp_path):
    run = _canvas(tmp_path)
    assert "a11y::image-alt" in write_report(run, ["fixes"])["fixes"].read_text()

    _triage(run, {"s001::image-alt::x": "wontfix", "s002::image-alt::x": "wontfix"})
    doc = build_document(run)
    assert all(f["audit"] != "a11y" for f in doc["top"])
    text = write_report(run, ["fixes"])["fixes"].read_text()
    assert "a11y::image-alt" not in text  # dismissed: neither open nor ticked as fixed
    state = json.loads((run / "reports" / "fixes.json").read_text())["tasks"]["a11y::image-alt"]
    assert state["dismissed"] == "wontfix" and "fixed_on" not in state


def test_mixed_triage_keeps_only_open_issues_as_where_to_fix(tmp_path):
    run = _canvas(tmp_path)
    _triage(run, {"s001::image-alt::x": "false_positive"})
    task = next(t for t in tasks_from(build_document(run)) if t["task"] == "a11y::image-alt")
    assert [w["id"] for w in task["where"]] == ["s002"]
    assert [e["screen"] for e in task["elements"]] == ["s002"]


def test_open_task_whose_audit_did_not_run_is_listed_as_not_rechecked(tmp_path):
    run = _canvas(tmp_path)
    write_report(run, ["fixes"])
    (run / "audits" / "store.json").unlink()
    text = write_report(run, ["fixes"])["fixes"].read_text()
    assert "## Not re-checked" in text and "`store::target`" in text


def test_another_tools_file_under_a_builtin_id_wins(tmp_path):
    theirs = {"schema": 1, "audit": "perf", "title": "Performance (Perfetto)", "tool": "perfetto 50",
              "status": "warn", "summary": "From a trace.", "findings": []}  # fmt: skip
    run = _canvas(tmp_path, extra={"perf": theirs})
    perf = next(s for s in build_document(run)["sections"] if s["id"] == "perf")
    assert perf["title"] == "Performance (Perfetto)"
    assert json.loads((run / "audits" / "perf.json").read_text())["tool"] == "perfetto 50"


def test_an_audit_holding_an_error_never_reads_as_passing():
    from appscreenmap.reports.builtin import perf

    data = {"run": {}, "screens": []}
    screens = [{"screen": f"s{i:03}", "title": f"S{i}", "median_ms": 300, "frames": 40, "janky_pct": 1,
                "rating": "good"} for i in range(19)]  # fmt: skip
    screens.append(
        {"screen": "s099", "title": "Web", "median_ms": 2700, "frames": 2, "janky_pct": 0, "rating": "slow"}
    )
    audit = perf(data, {"budget_ms": 1000, "screens": screens})
    assert audit["score"] == 95 and audit["status"] == "warn"


def test_the_same_crash_at_another_step_stays_the_same_task(tmp_path):
    run = _canvas(tmp_path)
    write_report(run, ["fixes"])
    data = json.loads((run / "map.json").read_text())
    data["events"] = [{"step": 9, "type": "crash", "screen": "s002", "detail": "NPE"}]
    (run / "map.json").write_text(json.dumps(data))
    text = write_report(run, ["fixes"])["fixes"].read_text()
    assert "Crash on Profile" in text and "## Fixed" not in text


def test_a11y_title_counts_only_issues_still_open(tmp_path):
    run = _canvas(tmp_path)
    _triage(run, {"s001::image-alt::x": "false_positive"})
    a11y = next(s for s in build_document(run)["sections"] if s["id"] == "a11y")
    assert a11y["findings"][0]["title"] == "Non-text Content (1 issue)"


def test_docx_keeps_sections_tables_links_and_thumbnails(tmp_path):
    from docx import Document

    run = _canvas(tmp_path)
    assert write_report(run, ["docx"])["docx"].is_file()  # the canvas's fake PNG is skipped, not fatal
    from PIL import Image

    Image.new("RGB", (40, 80), "white").save(run / "screens" / "s001.png")
    path = write_report(run, ["docx"])["docx"]
    d = Document(str(path))
    h1 = [p.text for p in d.paragraphs if p.style.name == "Heading 1"]
    assert h1[:2] == ["Summary", "Conditions"] and h1[-1] == "Appendix"
    assert "Store readiness" in h1
    assert d.tables[1].rows[0].cells[0].text == "Audit"  # the summary table, after the cover facts
    rels = [r.target_ref for r in d.part.rels.values()]
    assert "https://support.google.com/x" in rels  # sources are real hyperlinks
    assert any("image" in r.reltype for r in d.part.rels.values())


def test_pdf_reports_a_missing_browser_clearly(tmp_path, monkeypatch):
    import pytest

    from appscreenmap.reports import pdf

    run = _canvas(tmp_path)
    (run / "reports").mkdir()
    (run / "reports" / "report.pdf").write_bytes(b"%PDF last run")
    monkeypatch.setenv("APPSCREENMAP_CHROME", str(tmp_path / "no-such-chrome"))
    with pytest.raises(pdf.PdfError, match="doesn't exist"):
        write_report(run, ["pdf"])
    assert not (run / "reports" / "report.pdf").exists()  # never left beside this run's md / html


def test_pdf_reports_a_browser_that_cant_start(tmp_path, monkeypatch):
    import pytest

    from appscreenmap.reports import pdf

    fake = tmp_path / "chrome"
    fake.write_text("not a program")  # exists, not executable
    monkeypatch.setenv("APPSCREENMAP_CHROME", str(fake))
    with pytest.raises(pdf.PdfError, match="couldn't start"):
        write_report(_canvas(tmp_path), ["pdf"])


def _real_shots(run, ids):
    from PIL import Image

    for sid in ids:
        Image.new("RGB", (400, 800), (40, 40, 40)).save(run / "screens" / f"{sid}.png")


def test_report_shows_every_screen_and_outlined_evidence(tmp_path):
    from docx import Document

    issues = [
        {"key": f"{s}::image-alt::x", "screen": s, "rule": "image-alt", "sc": "1.1.1", "level": "A",
         "severity": "warning", "rid": "dev.alpha:id/logo", "label": "logo", "triage": "open",
         "bounds": [100, 200, 220, 320]}
        for s in ("s001", "s002")
    ]  # fmt: skip
    run = _canvas(tmp_path, a11y_issues=issues)
    data = json.loads((run / "map.json").read_text())
    data["screens"] = [{**sc, "screenshot": f"screens/{sc['id']}.png"} for sc in data["screens"]]
    (run / "map.json").write_text(json.dumps(data))
    _real_shots(run, ("s001", "s002"))
    out = write_report(run, ["md", "html", "docx"])

    doc = build_document(run)
    assert [g["id"] for g in doc["gallery"]] == ["s001", "s002"]
    assert doc["gallery"][0]["scenario"] == "First run"
    f = next(f for s in doc["sections"] for f in s["findings"] if f["id"] == "a11y::image-alt")
    assert [ev["id"] for ev in f["evidence"]] == ["s001", "s002"]

    html = out["html"].read_text()
    assert 'id="screens"' in html and html.count('class="evidence"') == 1 and "(outlined)" in html
    md = out["md"].read_text()
    assert "## Screens" in md and "**Evidence**" in md
    assert len(list((run / "reports" / "report-assets").glob("evidence-*.jpg"))) == 2
    # gallery (2) + a11y thumbnails (2) + crash thumbnail (1) + evidence (2)
    assert len(Document(str(out["docx"])).inline_shapes) == 7


def test_finding_names_screens_beyond_the_thumbnail_cap(tmp_path):
    from appscreenmap.reports.document import MAX_SHOTS

    ids = [f"s{n:03}" for n in range(1, MAX_SHOTS + 3)]
    issues = [
        {"key": f"{s}::image-alt::x", "screen": s, "rule": "image-alt", "sc": "1.1.1", "level": "A",
         "severity": "warning", "rid": "dev.alpha:id/logo", "label": "logo", "triage": "open"}
        for s in ids
    ]  # fmt: skip
    run = _canvas(tmp_path, a11y_issues=issues)
    data = json.loads((run / "map.json").read_text())
    data["screens"] = [_screen(s, s, thumb=f"screens/{s}.png") for s in ids]
    (run / "map.json").write_text(json.dumps(data))
    _real_shots(run, ids)
    html = write_report(run, ["html"])["html"].read_text()
    assert "+2 more" in html


def test_docx_skips_a_truncated_thumbnail(tmp_path):
    run = _canvas(tmp_path)
    data = json.loads((run / "map.json").read_text())
    data["screens"][0]["thumb"] = "screens/s001.jpg"
    (run / "map.json").write_text(json.dumps(data))
    (run / "screens" / "s001.jpg").write_bytes(b"\xff\xd8\xff\xe0" + b"\0" * 36)  # 40 bytes, cut short
    assert write_report(run, ["docx"])["docx"].is_file()


def test_pdf_prints_through_chrome_when_one_is_installed(tmp_path):
    import pytest

    from appscreenmap.reports import pdf

    if not pdf.find_chrome():
        pytest.skip("no Chrome or Chromium on this machine")
    path = write_report(_canvas(tmp_path), ["pdf"], paper="letter")["pdf"]
    assert path.read_bytes()[:5] == b"%PDF-"
    assert not (path.parent / ".report-print.html").exists()


def test_paper_palette_is_print_only():
    from appscreenmap.reports import htmldoc, pdf

    tokens = htmldoc.tokens_css()
    for t in pdf.PAPER_TOKENS:
        assert f"--print-{t}:" in tokens  # every remapped token exists in tokens.css
    css = pdf.print_css("a4")
    assert css.index("@media print{") < css.index("--bg:var(--print-bg)")
    assert "@media (prefers-color-scheme" not in htmldoc.CSS + tokens  # screens stay dark


def test_evidence_skips_bounds_it_cannot_draw(tmp_path):
    from PIL import Image

    from appscreenmap.reports.document import has_area
    from appscreenmap.reports.evidence import crop

    Image.new("RGB", (400, 800)).save(tmp_path / "s.png")
    for b in (
        [-500, -500, -400, -400],
        [-1080, 200, -10, 400],
        ["a", "b", "c", "d"],
        [None, 1, 2, 3],
        [0, 0, 0, 0],
        [0, 0, float("inf"), 10],
    ):
        assert crop(tmp_path, "s.png", b) is None, b
        assert not has_area(b), b
    assert crop(tmp_path, "s.png", [10, 10, 100, 100])

    issues = [
        {"key": f"s001::image-alt::{n}", "screen": "s001", "rule": "image-alt", "sc": "1.1.1", "level": "A",
         "severity": "warning", "rid": f"dev.alpha:id/x{n}", "label": f"x{n}", "triage": "open", "bounds": b}
        for n, b in enumerate(([-500, -500, -400, -400], ["a", "b", "c", "d"], [0, 0, 0, 0]))
    ]  # fmt: skip
    run = _canvas(tmp_path / "c", a11y_issues=issues)
    data = json.loads((run / "map.json").read_text())
    data["screens"] = [{**sc, "screenshot": f"screens/{sc['id']}.png"} for sc in data["screens"]]
    (run / "map.json").write_text(json.dumps(data))
    _real_shots(run, ("s001", "s002"))
    out = write_report(run, ["md", "html", "docx"])
    assert "Evidence" not in out["html"].read_text().split('id="a11y"')[1]


def test_print_keeps_borders_off_the_printable_edge():
    from appscreenmap.reports.pdf import page_margin, print_css

    css = print_css("a4")
    assert "main{max-width:none;padding:0 var(--space-1)}" in css  # Chrome clips a border on the edge
    assert f"margin:{page_margin()}" in css and page_margin() == "14mm 12mm"
